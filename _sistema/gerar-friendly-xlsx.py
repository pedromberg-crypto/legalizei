# -*- coding: utf-8 -*-
"""cnae-friendly.json -> .xlsx, formatado pra REVISÃO (não pra importar).

O oficial e o amigável ficam lado a lado, porque a pergunta que o Pedro faz ao
revisar é "o de baixo diz a mesma coisa que o de cima, de um jeito que a pessoa
entende?". Uma coluna vazia de OK/ajuste no fim, pra ele marcar direto.
"""
import json, io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

dados = json.load(io.open('pesquisa/cnae-matriz/cnae-friendly.json', encoding='utf-8'))

wb = Workbook()
ws = wb.active
ws.title = "CNAEs atendidos (ME)"

INK = "1B1E24"
CORAL = "E85D2F"
PAPEL = "FBF7F2"
LINHA = "E6E0D8"

cab = ["#", "code", "official_description (IBGE)", "friendly_title", "friendly_description", "OK? / ajuste"]
ws.append(cab)

for i, d in enumerate(dados, start=1):
    ws.append([i, d["code"], d["official_description"], d["friendly_title"],
               d["friendly_description"], ""])

# ── cabeçalho
for c in range(1, len(cab) + 1):
    cel = ws.cell(row=1, column=c)
    cel.font = Font(bold=True, color="FFFFFF", size=11)
    cel.fill = PatternFill("solid", fgColor=INK)
    cel.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
ws.row_dimensions[1].height = 30

larguras = [5, 11, 52, 38, 62, 22]
for i, w in enumerate(larguras, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

borda = Border(bottom=Side(style="thin", color=LINHA))
for r in range(2, len(dados) + 2):
    for c in range(1, len(cab) + 1):
        cel = ws.cell(row=r, column=c)
        cel.alignment = Alignment(vertical="top", wrap_text=(c in (3, 4, 5, 6)))
        cel.border = borda
        if c == 2:
            cel.font = Font(name="Consolas", size=10)
        if c == 3:
            cel.font = Font(size=9, color="6B6B6B")   # o oficial é referência, não protagonista
        if c == 4:
            cel.font = Font(bold=True, size=11, color=CORAL)
    # zebra, pra o olho não pular de linha numa tabela de 87
    if r % 2 == 0:
        for c in range(1, len(cab) + 1):
            ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor=PAPEL)
    ws.row_dimensions[r].height = 34

ws.freeze_panes = "C2"          # code e # ficam à vista ao rolar
ws.auto_filter.ref = f"A1:F{len(dados) + 1}"

wb.save('pesquisa/cnae-matriz/cnae-friendly.xlsx')
print('gerado: pesquisa/cnae-matriz/cnae-friendly.xlsx |', len(dados), 'linhas')
